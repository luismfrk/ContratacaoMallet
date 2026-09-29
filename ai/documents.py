from __future__ import annotations

from io import BytesIO
from pathlib import Path

import pdfplumber
from docx import Document

MAX_FILE_BYTES = 5 * 1024 * 1024
MAX_EXTRACTED_CHARS = 24_000
ALLOWED_EXTENSIONS = {".pdf", ".docx", ".txt"}


def extract_example_text(filename: str, content: bytes) -> str:
    extension = Path(filename or "").suffix.lower()
    if extension not in ALLOWED_EXTENSIONS:
        raise ValueError("Envie um arquivo PDF, DOCX ou TXT.")
    if not content:
        raise ValueError("O arquivo enviado está vazio.")
    if len(content) > MAX_FILE_BYTES:
        raise ValueError("O arquivo deve ter no máximo 5 MB.")

    try:
        if extension == ".txt":
            text = content.decode("utf-8-sig")
        elif extension == ".docx":
            document = Document(BytesIO(content))
            parts = [paragraph.text for paragraph in document.paragraphs]
            parts.extend(
                " | ".join(cell.text for cell in row.cells)
                for table in document.tables for row in table.rows
            )
            text = "\n".join(parts)
        else:
            with pdfplumber.open(BytesIO(content)) as pdf:
                text = "\n".join(page.extract_text() or "" for page in pdf.pages)
    except Exception as exc:
        raise ValueError("Não foi possível ler o arquivo. Verifique se ele está íntegro.") from exc

    text = text.strip()
    if not text:
        raise ValueError("Não foi encontrado texto legível no arquivo.")
    return text[:MAX_EXTRACTED_CHARS]
