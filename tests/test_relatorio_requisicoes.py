from io import BytesIO
import unittest

from openpyxl import load_workbook

from relatorio_requisicoes import gerar_relatorio_requisicoes


class TestRelatorioRequisicoes(unittest.TestCase):
    def test_relatorio_vazio_nao_cria_referencia_circular(self):
        conteudo, _ = gerar_relatorio_requisicoes([], 2026, 8, "")
        ws = load_workbook(BytesIO(conteudo), data_only=False).active

        self.assertEqual(ws["C5"].value, 0)
        self.assertEqual(ws["D5"].value, 0)
        self.assertEqual(ws["E5"].value, 0)

    def test_total_soma_apenas_as_linhas_de_dados(self):
        linhas = [{
            "placa": "ABC-1234",
            "requisicoes": 1,
            "materiais": 100,
            "servicos": 50,
            "total": 150,
        }]
        conteudo, _ = gerar_relatorio_requisicoes(linhas, 2026, 8, "Educação")
        ws = load_workbook(BytesIO(conteudo), data_only=False).active

        self.assertEqual(ws["C6"].value, "=SUM(C5:C5)")
        self.assertEqual(ws["D6"].value, "=SUM(D5:D5)")
        self.assertEqual(ws["E6"].value, "=SUM(E5:E5)")


if __name__ == "__main__":
    unittest.main()
